import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font

# --- CONFIGURATION ---
TEMPLATE_FILE = 'template.xlsx'
FAKULTAS = 20
PRODI = 138
TAHUN = 2025
SEMESTER = 2
SEMESTER_KELAS = 2
TARGET_KELAS = 'D'


OUTPUT_FILE = f"JADWAL KELAS {TARGET_KELAS} {PRODI} SEM {SEMESTER_KELAS} {TAHUN}.xlsx"

# Define pastel colors for lecturers (Hex codes)
# We will cycle through these if there are many lecturers
COLORS = [
    "FFD1DC", "FFDFD3", "FFFFD1", "E2F0CB", "B5EAD7", "C7CEEA", "FF9AA2", "FFB7B2", "FFDAC1", "E2F0CB", "B5EAD7", "C7CEEA", "F0E68C", "E6E6FA", "D8BFD8", "FFC0CB", "98FB98", "AFEEEE"
]

def get_schedule_html():
    print("Connecting to UNS server...")
    s = requests.Session()

    # Step 1
    r = s.get("https://jadwal.uns.ac.id/jadwal")
    soup = BeautifulSoup(r.text, "html.parser")
    try:
        csrf = soup.find("input", {"name": "_csrf-frontend"})["value"]
    except TypeError:
        print("Error: Could not find CSRF token. The site might be down or changed.")
        return None
    
    # Step 2: Request the schedule
    payload = {
        "_csrf-frontend": csrf,
        "FAKULTAS[IDFAKULTAS]": FAKULTAS,
        "PRODI[IDPRODI]": PRODI,
        "TAS[TAHUNAJAR]": TAHUN,
        "TAS[IDSEMESTER]": SEMESTER,
        "SEMESTER[IDSEMESTER]": SEMESTER_KELAS,
    }

    print("Fetching schedule data...")
    jadwal = s.post("https://jadwal.uns.ac.id/jadwal/cetak", data=payload)
    return jadwal.text

def parse_html_to_data(html_content, target_class):
    soup = BeautifulSoup(html_content, "html.parser")
    tables = soup.find_all("table", class_="table-bordered")
    
    schedule_data = []

    print(f"Parsing data for Kelas {target_class}...")
    
    for table in tables:
        rows = table.find("tbody").find_all("tr")
        for row in rows:
            cols = row.find_all("td")
            if not cols:
                continue
            
            # Map HTML columns based on your provided HTML structure:
            # 0: No, 1: Hari, 2: Sesi, 3: Jam, 4: Makul, 5: Kode, 
            # 6: Sem, 7: SKS, 8: Dosen, 9: Ruang, 10: Kelas
            
            kelas_raw = cols[10].text.strip()
            
            # Filter by class
            if kelas_raw != target_class:
                continue

            entry = {
                "hari": cols[1].text.strip(),
                "sesi": int(cols[2].text.strip()),
                "makul": cols[4].text.strip(),
                "dosen": cols[8].text.strip(),
                "ruang": cols[9].text.strip(),
            }
            schedule_data.append(entry)
            
    return schedule_data

def map_and_write_excel(data, template_path, output_path):
    print("Opening Excel template...")
    try:
        wb = openpyxl.load_workbook(template_path)
        ws = wb.active # Assumes data goes on the first sheet
    except FileNotFoundError:
        print(f"Error: {template_path} not found.")
        return

    # Map 'Hari' string to Excel Column Index (Based on your CSV)
    # Senin=B(2), Selasa=C(3), Rabu=D(4), Kamis=E(5), Jumat=F(6)
    day_map = {
        "Senin": 2, "Selasa": 3, "Rabu": 4, "Kamis": 5, "Jumat": 6
    }

    # Map 'Sesi' integer to Excel Row Index
    # Sesi 1 is Row 3, Sesi 6 is Row 9, etc.
    session_row_map = {
        1: 3, 2: 4, 3: 5, 4: 6, 5: 7,       # Morning block
        6: 9, 7: 10,                        # After 1st break
        8: 12, 9: 13,                       # After 2nd break
        10: 15                              # Evening block
    }

    # Create unique colors for each lecturer
    unique_lecturers = sorted(list(set(item['dosen'] for item in data)))
    lecturer_colors = {}
    
    print("Assigning colors to lecturers...")
    # Write Legend (Right side of Excel)
    legend_start_row = 3
    legend_col_dosen = 8 # Column H
    legend_col_color = 9 # Column I

    for idx, dosen in enumerate(unique_lecturers):
        # Pick a color (cycle if we run out)
        hex_color = COLORS[idx % len(COLORS)]
        fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        lecturer_colors[dosen] = fill

        # Write to Legend columns
        ws.cell(row=legend_start_row + idx, column=legend_col_dosen).value = dosen
        color_cell = ws.cell(row=legend_start_row + idx, column=legend_col_color)
        color_cell.fill = fill

    print("Filling schedule grid...")
    for item in data:
        day_str = item['hari']
        session_int = item['sesi']
        
        if day_str not in day_map or session_int not in session_row_map:
            print(f"Warning: Data out of bounds ({day_str}, Sesi {session_int})")
            continue

        target_col = day_map[day_str]
        target_row = session_row_map[session_int]
        
        # Prepare Cell Content
        cell_content = f"{item['makul']}\n({item['ruang']})"
        
        cell = ws.cell(row=target_row, column=target_col)
        cell.value = cell_content
        
        # Apply Styling
        cell.fill = lecturer_colors[item['dosen']]
        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        cell.font = Font(name='Times New Roman', size=12)

    wb.save(output_path)
    print(f"Success! Schedule saved to {output_path}")

if __name__ == "__main__":
    # 1. Get HTML (Or use a local file for testing if you save 'output.txt')
    html_source = get_schedule_html()

    if html_source:
        # 2. Parse Data
        parsed_data = parse_html_to_data(html_source, TARGET_KELAS)
        
        # 3. Write to Excel
        if parsed_data:
            map_and_write_excel(parsed_data, TEMPLATE_FILE, OUTPUT_FILE)
        else:
            print(f"No data found for Class {TARGET_KELAS}. Check the HTML source.")